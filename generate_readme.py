import openpyxl
import os
import re

script_dir = os.path.dirname(os.path.abspath(__file__))
excel_path = os.path.join(script_dir, "DSA_Complete_Roadmap_1_to_Interview.xlsx")
readme_path = os.path.join(script_dir, "README.md")
svg_path = os.path.join(script_dir, "progress.svg")

if not os.path.exists(excel_path):
    print("Excel file not found at:", excel_path)
    exit(1)

wb = openpyxl.load_workbook(excel_path, data_only=True)
sheet = wb["Master Roadmap"]

# Helper to clean name for filename
def clean_filename(problem_name):
    name = re.sub(r'\([^)]*\)', '', problem_name)
    words = re.split(r'[\s\-/\+\\\._]+', name)
    clean_words = []
    for w in words:
        if w:
            w_clean = re.sub(r'[^a-zA-Z0-9]', '', w)
            if w_clean:
                clean_words.append(w_clean.capitalize())
    return "".join(clean_words)

# Mapping of custom solved files
solved_mapping = {
    1: "23Celibrity.cpp",
    2: "24MajorityElement.cpp",
    3: "7TwoSum.cpp",
    6: "14palandrom.cpp",
    20: "17reverselist.cpp",
    52: "11rotateArray.cpp",
    53: "6mergeSortedArrays.cpp",
    92: "8maximumSubarray.cpp"
}

# Detailed problem descriptions for on-the-go reading (walking the dog)
problem_details = {
    9: {
        "statement": "Given an array `nums`. We define a running sum of an array as `runningSum[i] = sum(nums[0]…nums[i])`.\n\nReturn the running sum of `nums`.",
        "examples": [
            "**Input:** `nums = [1, 2, 3, 4]`\n**Output:** `[1, 3, 6, 10]`\n**Explanation:** `[1, 1+2, 1+2+3, 1+2+3+4]` -> `[1, 3, 6, 10]`",
            "**Input:** `nums = [1, 1, 1, 1, 1]`\n**Output:** `[1, 2, 3, 4, 5]`"
        ],
        "constraints": "- $1 \\le \\text{nums.length} \\le 1000$\n- $-10^6 \\le \\text{nums}[i] \\le 10^6$",
        "target_complexity": "Time: $O(N)$ | Space: $O(1)$ extra space (modify in-place)"
    },
    10: {
        "statement": "Given an array of integers `nums` and an integer `k`, return *the total number of subarrays whose sum equals to `k`*.\n\nA subarray is a contiguous non-empty sequence of elements within an array.",
        "examples": [
            "**Input:** `nums = [1, 1, 1], k = 2`\n**Output:** `2` (subarrays `[1,1]` at indices 0-1 and 1-2)",
            "**Input:** `nums = [1, 2, 3], k = 3`\n**Output:** `2` (subarrays `[1,2]` and `[3]`)"
        ],
        "constraints": "- $1 \\le \\text{nums.length} \\le 2 \\times 10^4$\n- $-1000 \\le \\text{nums}[i] \\le 1000$\n- $-10^7 \\le k \\le 10^7$",
        "target_complexity": "Time: $O(N)$ | Space: $O(N)$ with Prefix Sum + Hash Map"
    },
    11: {
        "statement": "Given an array of positive integers `nums` and a positive integer `target`, return the **minimal length** of a subarray whose sum is greater than or equal to `target`. If there is no such subarray, return `0` instead.",
        "examples": [
            "**Input:** `target = 7, nums = [2, 3, 1, 2, 4, 3]`\n**Output:** `2`\n**Explanation:** The subarray `[4, 3]` has the minimal length under the problem constraint.",
            "**Input:** `target = 4, nums = [1, 4, 4]`\n**Output:** `1`"
        ],
        "constraints": "- $1 \\le \\text{target} \\le 10^9$\n- $1 \\le \\text{nums.length} \\le 10^5$\n- $1 \\le \\text{nums}[i] \\le 10^4$",
        "target_complexity": "Time: $O(N)$ | Space: $O(1)$ with Sliding Window"
    },
    12: {
        "statement": "Given a string `s`, find the length of the **longest substring** without duplicate characters.",
        "examples": [
            "**Input:** `s = \"abcabcbb\"`\n**Output:** `3` (The answer is `\"abc\"`, with the length of 3)",
            "**Input:** `s = \"bbbbb\"`\n**Output:** `1` (The answer is `\"b\"`, with length of 1)"
        ],
        "constraints": "- $0 \\le \\text{s.length} \\le 5 \\times 10^4$\n- `s` consists of English letters, digits, symbols and spaces.",
        "target_complexity": "Time: $O(N)$ | Space: $O(\\min(N, M))$ with Sliding Window + Hash Set"
    },
    13: {
        "statement": "Given two strings `s1` and `s2`, return `true` if `s2` contains a **permutation** of `s1`, or `false` otherwise.\n\nIn other words, return `true` if one of `s1`'s permutations is the substring of `s2`.",
        "examples": [
            "**Input:** `s1 = \"ab\", s2 = \"eidbaooo\"`\n**Output:** `true` (s2 contains one permutation of s1: `\"ba\"`)",
            "**Input:** `s1 = \"ab\", s2 = \"eidboaoo\"`\n**Output:** `false`"
        ],
        "constraints": "- $1 \\le \\text{s1.length}, \\text{s2.length} \\le 10^4$\n- `s1` and `s2` consist of lowercase English letters.",
        "target_complexity": "Time: $O(N)$ | Space: $O(1)$ (array of 26)"
    },
    14: {
        "statement": "Given a string `s` containing just the characters `'('`, `')'`, `'{'`, `'}'`, `'['` and `']'`, determine if the input string is valid.\n\nAn input string is valid if:\n1. Open brackets must be closed by the same type of brackets.\n2. Open brackets must be closed in the correct order.\n3. Every close bracket has a corresponding open bracket of the same type.",
        "examples": [
            "**Input:** `s = \"()[]{}\"`\n**Output:** `true`",
            "**Input:** `s = \"(]\"`\n**Output:** `false`"
        ],
        "constraints": "- $1 \\le \\text{s.length} \\le 10^4$\n- `s` consists of parentheses only `'()[]{}'`.",
        "target_complexity": "Time: $O(N)$ | Space: $O(N)$ with Stack"
    }
}

problems = []
solved_count = 0

for r in range(2, sheet.max_row + 1):
    no_val = sheet.cell(row=r, column=1).value
    level = sheet.cell(row=r, column=2).value
    pattern = sheet.cell(row=r, column=3).value
    problem_name = sheet.cell(row=r, column=4).value
    difficulty = sheet.cell(row=r, column=5).value
    why_skill = sheet.cell(row=r, column=6).value
    status = sheet.cell(row=r, column=7).value
    
    if not problem_name:
        continue
        
    try:
        no_val = int(no_val)
    except (TypeError, ValueError):
        continue
        
    is_solved = (status == "Solved")
    if is_solved:
        solved_count += 1
        
    # Get local filename link
    if no_val in solved_mapping:
        local_file = solved_mapping[no_val]
    else:
        safe_name = clean_filename(problem_name)
        local_file = f"No_{no_val:03d}_{safe_name}.cpp"
        
    problems.append({
        'no': no_val,
        'level': level,
        'pattern': pattern,
        'name': problem_name,
        'difficulty': difficulty,
        'why': why_skill,
        'solved': is_solved,
        'file': local_file
    })

total_problems = len(problems)
pct = (solved_count / total_problems) * 100 if total_problems else 0
fill_width = 400 * (solved_count / total_problems) if total_problems else 0

# Count by level
level_stats = {}
for p in problems:
    lvl = p['level']
    if lvl not in level_stats:
        level_stats[lvl] = {'total': 0, 'solved': 0}
    level_stats[lvl]['total'] += 1
    if p['solved']:
        level_stats[lvl]['solved'] += 1

# Find first unsolved problem (Current Target)
current_target = None
for p in problems:
    if not p['solved']:
        current_target = p
        break

# Generate progress.svg
svg_content = f"""<svg width="400" height="60" viewBox="0 0 400 60" fill="none" xmlns="http://www.w3.org/2000/svg">
  <style>
    .track {{
      fill: #e1e4e8;
    }}
    .fill {{
      fill: url(#grad);
      animation: load 2s cubic-bezier(0.4, 0, 0.2, 1) forwards;
      width: 0px;
    }}
    .text {{
      font-family: -apple-system, BlinkMacSystemFont, "Segoe UI", Helvetica, Arial, sans-serif;
      font-size: 14px;
      font-weight: 600;
      fill: #24292e;
      opacity: 0;
      animation: fadeIn 0.5s ease-out 1s forwards;
    }}
    @keyframes load {{
      to {{ width: {fill_width:.1f}px; }}
    }}
    @keyframes fadeIn {{
      to {{ opacity: 1; }}
    }}
    @media (prefers-color-scheme: dark) {{
      .track {{
        fill: #21262d;
        stroke: #30363d;
      }}
      .text {{
        fill: #c9d1d9;
      }}
    }}
  </style>
  <defs>
    <linearGradient id="grad" x1="0%" y1="0%" x2="100%" y2="0%">
      <stop offset="0%" stop-color="#4facfe" />
      <stop offset="100%" stop-color="#00f2fe" />
    </linearGradient>
  </defs>
  
  <text x="0" y="18" class="text" style="animation: fadeIn 0.5s ease-out forwards; opacity: 1;">Roadmap Completion</text>
  <text x="360" y="18" class="text">{pct:.1f}%</text>

  <rect x="0" y="30" width="400" height="16" rx="8" class="track" />
  <rect x="0" y="30" height="16" rx="8" class="fill" />
</svg>
"""

with open(svg_path, "w", encoding="utf-8") as f:
    f.write(svg_content)

# Target Problem Card Generator
target_section = ""
if current_target:
    c_no = current_target['no']
    c_name = current_target['name']
    c_diff = current_target['difficulty']
    c_pattern = current_target['pattern']
    c_file = current_target['file']
    c_slug = c_name.lower().replace(' ', '-').replace('(', '').replace(')', '').replace("'", "").replace(",", "")
    c_link = f"https://leetcode.com/problems/{c_slug}/"
    diff_badge = "🟢 Easy" if c_diff == "Easy" else ("🟡 Medium" if c_diff == "Medium" else "🔴 Hard")

    details = problem_details.get(c_no, {})
    statement = details.get("statement", f"Solve **{c_name}** on [LeetCode]({c_link}).")
    examples_md = "\n\n".join(details.get("examples", []))
    constraints = details.get("constraints", "Standard LeetCode constraints apply.")
    target_comp = details.get("target_complexity", "Aim for the most optimal Time and Space complexity.")

    target_section = f"""---

## 🎯 Current Target Problem (Solve on the Go)

> **📱 Mobile Friendly Card**: Read the problem, test cases, and constraints right here on your phone while walking or commuting!

### **No. {c_no}: [{c_name}]({c_link})**
* **Level:** `{current_target['level']}` | **Difficulty:** {diff_badge} | **Pattern:** `{c_pattern}`
* **Local Code:** [`{c_file}`](./{c_file})

#### 📝 Problem Statement
{statement}

#### 💡 Examples
{examples_md if examples_md else "*See LeetCode for examples.*"}

#### ⚙️ Constraints & Target
{constraints}
* **Optimal Complexity Goal:** {target_comp}

---
"""

# Generate README Content with badges, active problem card, and animated SVG
content = f"""# 🎯 ZTM DSA: 0 to Hero Placement Roadmap

Welcome to my personal Data Structures and Algorithms (DSA) workspace! This repository contains my implementations of foundational data structures from Andrei Neagoie's Zero To Mastery (ZTM) course, along with a curated **150-problem LeetCode roadmap** aimed at cracking product-based company placement drives.

<p align="left">
  <img src="https://img.shields.io/badge/LeetCode-{solved_count}%20%2F%20150%20Solved-blue?style=for-the-badge&logo=leetcode" alt="Leetcode Solved">
  <img src="https://img.shields.io/badge/Language-C%2B%2B-orange?style=for-the-badge&logo=c%2B%2B" alt="Language">
  <img src="https://img.shields.io/badge/Placement-Prep-success?style=for-the-badge" alt="Placement Prep">
  <img src="https://img.shields.io/badge/Status-Grinding-red?style=for-the-badge" alt="Status">
</p>

---

## 📊 Progress Dashboard

Here is my interactive progress tracker. It features a custom animated bar that automatically updates whenever a problem is marked as solved!

<p align="left">
  <img src="./progress.svg" alt="Roadmap Progress" width="400">
</p>

| Difficulty Level | Problems Solved | Progress Percentage | Status |
| :--- | :---: | :---: | :---: |
"""

for lvl, stats in level_stats.items():
    l_pct = (stats['solved'] / stats['total']) * 100
    indicator = "🟢 Active" if stats['solved'] > 0 else "⚪ Not Started"
    if stats['solved'] == stats['total']:
        indicator = "🏆 Completed"
    content += f"| **{lvl}** | {stats['solved']} / {stats['total']} | {l_pct:.1f}% | {indicator} |\n"

content += f"""| **TOTAL** | **{solved_count} / {total_problems}** | **{pct:.1f}%** | **🧠 Grind Mode** |

{target_section}

## 🛠️ Tech Stack & Structure
* **Language:** C++ (using modern Standard Template Library - STL)
* **Compiler:** `g++` (C++17/C++20 standard)
* **Custom Codebases:** Core data structures (BST, Graphs, LinkedLists, HashTables) coded from scratch to build core pointer memory concepts.
* **Problem Set:** Curated by difficulty levels and algorithmic patterns (Sliding Window, Monotonic Stacks, DP, Dijkstra, Segment Trees).

---

## 📂 Detailed Roadmap

Below is the complete list of 150 problems categorized by level. Each problem links to its LeetCode description page and its local C++ source/template file.

"""

levels_ordered = ["Level 1", "Level 2", "Level 3", "Interview Mode"]

for lvl in levels_ordered:
    lvl_problems = [p for p in problems if p['level'] == lvl]
    stats = level_stats.get(lvl, {'total': 0, 'solved': 0})
    
    content += f"""
<details>
<summary><b>🔥 {lvl} ({stats['solved']}/{stats['total']} Solved)</b></summary>
<br>

| No. | Problem | Algorithmic Pattern | Difficulty | Local C++ Code | Status |
| :---: | :--- | :--- | :---: | :--- | :---: |
"""
    for p in lvl_problems:
        status_icon = "✅ Solved" if p['solved'] else "⏳ Not Started"
        diff_color = "🟢 Easy" if p['difficulty'] == "Easy" else ("🟡 Medium" if p['difficulty'] == "Medium" else "🔴 Hard")
        
        # Leetcode slug
        lc_slug = p['name'].lower().replace(' ', '-').replace('(', '').replace(')', '').replace("'", "").replace(",", "")
        lc_link = f"https://leetcode.com/problems/{lc_slug}/"
        
        content += f"| {p['no']} | [{p['name']}]({lc_link}) | {p['pattern']} | {diff_color} | [`{p['file']}`](./{p['file']}) | {status_icon} |\n"
        
    content += "\n</details>\n"

content += """
---

## 🚀 How to Compile and Run (C++)

To compile and run any C++ solution locally, use a terminal and run the following commands:

```bash
# Compile
g++ -std=c++17 -O3 <Filename>.cpp -o output.exe

# Execute (Windows)
./output.exe
```

---

*“The only way to learn a new programming language is by writing programs in it.” – Dennis Ritchie. Keep grinding!* 💻🔥
"""

with open(readme_path, "w", encoding="utf-8") as f:
    f.write(content)

print(f"SUCCESS: Re-generated README.md with Active Target Card and progress.svg! (Solved: {solved_count}/{total_problems})")
